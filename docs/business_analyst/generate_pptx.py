import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_ba_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    TOTAL_SLIDES = 16

    # Professional Executive Color Palette
    C_NAVY = RGBColor(0, 47, 94)        # #002f5e - Navy Port Authority
    C_BLUE = RGBColor(1, 112, 185)      # #0170b9 - Primary Blue
    C_LIGHT_BG = RGBColor(248, 250, 252)# #f8fafc - Slate 50
    C_CARD_BG = RGBColor(255, 255, 255) # #ffffff - Card White
    C_BORDER = RGBColor(226, 232, 240)  # #e2e8f0 - Border Gray
    C_DARK = RGBColor(15, 23, 42)       # #0f172a - Slate 900
    C_TEXT = RGBColor(30, 41, 59)       # #1e293b - Slate 800
    C_MUTED = RGBColor(100, 116, 139)   # #64748b - Slate 500
    C_EMERALD = RGBColor(16, 185, 129)  # #10b981 - Success Green
    C_AMBER = RGBColor(245, 158, 11)    # #f59e0b - Accent Amber
    C_INDIGO = RGBColor(99, 102, 241)   # #6366f1 - Tech Indigo
    C_PURPLE = RGBColor(147, 51, 234)   # #9333ea - Purple
    C_WHITE = RGBColor(255, 255, 255)
    C_GREEN_BG = RGBColor(236, 253, 245)# #ecfdf5
    C_BLUE_BG = RGBColor(239, 246, 255) # #eff6ff
    C_AMBER_BG = RGBColor(254, 243, 199) # #fef3c7
    C_RED_BG = RGBColor(254, 242, 242)  # #fef2f2
    C_RED = RGBColor(239, 68, 68)

    base_dir = r"C:\xampp\htdocs\Conclusion_Intermodal_Dry_Port\docs\business_analyst"

    def add_header(slide, title_text, category_text="CONCLUSION SUPPLY CHAIN CONSULTANT — STRATEGIC BUSINESS BLUEPRINT", web_page=None, web_url=None):
        box_width = Inches(8.5) if web_page else Inches(11.7)
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.28), box_width, Inches(1.15))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(8.5)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(16.5) if web_page else Pt(18.5)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p1.space_before = Pt(2)

        if web_page:
            badge_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.2), Inches(0.26), Inches(3.3), Inches(1.02))
            badge_card.fill.solid()
            badge_card.fill.fore_color.rgb = C_BLUE_BG
            badge_card.line.color.rgb = C_BLUE
            badge_card.line.width = Pt(1)
            btf = badge_card.text_frame
            btf.word_wrap = True
            btf.margin_left = Inches(0.12)
            btf.margin_right = Inches(0.12)
            btf.margin_top = Inches(0.08)
            btf.margin_bottom = Inches(0.08)
            
            bp0 = btf.paragraphs[0]
            bp0.text = "🌐 HALAMAN WEB / LIVE DEMO"
            bp0.font.size = Pt(7.5)
            bp0.font.bold = True
            bp0.font.color.rgb = C_BLUE
            
            bp1 = btf.add_paragraph()
            bp1.text = web_page
            bp1.font.size = Pt(9.5)
            bp1.font.bold = True
            bp1.font.color.rgb = C_NAVY
            bp1.space_before = Pt(1)
            
            if web_url:
                bp2 = btf.add_paragraph()
                bp2.text = web_url
                bp2.font.size = Pt(7)
                bp2.font.italic = True
                bp2.font.color.rgb = C_MUTED
                bp2.space_before = Pt(1)

    def add_footer(slide, current_slide):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35))
        tf = footer_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Conclusion Supply Chain Consultant  •  Naufal Andika Heditya & Zulfikar J. Fatah  •  Topik 7: CIDP 35 Ha                    Slide {current_slide} of {TOTAL_SLIDES}"
        p.font.size = Pt(8.5)
        p.font.color.rgb = C_MUTED

    def create_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def add_demo_banner(slide, text="[LIVE DEMO]: Buka Halaman Web", left=Inches(0.8), top=Inches(6.40), width=Inches(11.7), height=Inches(0.48)):
        b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        b.fill.solid()
        b.fill.fore_color.rgb = C_AMBER_BG
        b.line.color.rgb = C_AMBER
        b.line.width = Pt(1)
        tf = b.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"👉  {text}"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(180, 83, 9)
        p.alignment = PP_ALIGN.LEFT

    # =========================================================================
    # SLIDE 1: COVER KONSULTAN
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY
    bg1.line.fill.background()

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(0.12), Inches(6.4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_BLUE
    bar.line.fill.background()

    tbox1 = s1.shapes.add_textbox(Inches(1.1), Inches(0.45), Inches(11.4), Inches(2.2))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "CONCLUSION SUPPLY CHAIN CONSULTANT — STRATEGIC BUSINESS BLUEPRINT"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    p = tf1.add_paragraph()
    p.text = "Business Analysis & Alur Proses Bisnis End-to-End:\nPemetaan Fitur Software Logistik CIDP 35 Ha"
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_before = Pt(4)

    p = tf1.add_paragraph()
    p.text = "STANDAR SILABUS LENGKAP: 14 BAB PEMETAAN FITUR, BUSINESS RULES, EQUIPMENT ASSIGNMENT, FLOWCHART RESMI, & FSD"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_AMBER
    p.space_before = Pt(6)

    team_top = Inches(2.85)
    col_w = Inches(5.6)
    card_h = Inches(3.25)

    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), team_top, col_w, card_h)
    c1.fill.solid()
    c1.fill.fore_color.rgb = RGBColor(15, 33, 64)
    c1.line.color.rgb = RGBColor(30, 58, 110)
    c1.line.width = Pt(1)
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "TIM ADVISORY BUSINESS ANALYSIS & QA"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p = tf_c1.add_paragraph()
    p.text = "• Naufal Andika Heditya (24D507001012) — Lead Business Analyst & QA\n  Fokus: End-to-End Decision Flowcharts, Business Rules Timbangan, Equipment Allocation Matrix, Stacking Rules 40ft/20ft, & Spesifikasi FSD.\n\n• Zulfikar Jafarudin Fatah (24D507001001) — Lead System Architect\n  Fokus: Arsitektur Eksekutif Web YMS, Integrasi 5 Modul (WMS, TMS, WCS, YMS, ERP), & Database MySQL."
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(6)

    c2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), team_top, col_w, card_h)
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(15, 33, 64)
    c2.line.color.rgb = RGBColor(30, 58, 110)
    c2.line.width = Pt(1)
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "DEWAN KONSULTAN PENDUKUNG (CONCLUSION CONSULTANT)"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p = tf_c2.add_paragraph()
    p.text = "• Armansyah Muchtarrom (24D507001010) — Hardware Specialist\n  (Load Cell Timbangan 80T, Kamera ANPR/OCR, Palang Otomatis, & Sensor Telemetri)\n\n• Afriansayah Ayubi (24D507001026) — Software & ERP Specialist\n  (CFS Stripping/Stuffing WMS, Terminal Tariff Engine Odoo, & Billing)\n\n• Juan Gamaliel (24D507001016) — Data Integration Specialist\n  (REST API Gateway, Stacking Hierarchy 4D, & Exception Handling Engine)"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(6)

    fbox = s1.shapes.add_textbox(Inches(1.1), Inches(6.25), Inches(11.4), Inches(0.9))
    ftf = fbox.text_frame
    ftf.word_wrap = True
    p = ftf.paragraphs[0]
    p.text = "Klien / Evaluator Akademik: Dr. Tigor Franky, S.T., M.T.  •  Mata Kuliah: Teknologi dan Perangkat Lunak Logistik"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(203, 213, 225)
    p = ftf.add_paragraph()
    p.text = "Program Studi S1 Manajemen Logistik  •  Institut Transportasi dan Logistik (ITL) Trisakti  •  Oktober 2026"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p = ftf.add_paragraph()
    p.text = "🌐 Live Web Application Portal: http://localhost/Conclusion_Intermodal_Dry_Port/dashboard.php"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    # =========================================================================
    # SLIDE 2: PARADIGMA BUSINESS ANALYST: DEVELOPER VS BA & METODOLOGI (BAB 1)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Paradigma Business Analyst: Developer vs BA & Metodologi 3-Langkah (Bab 1)", web_page="dashboard.php?page=beranda", web_url="http://localhost/.../dashboard.php")
    add_footer(s2, 2)

    card_w = Inches(3.7)
    card_h = Inches(4.55)

    create_card(s2, Inches(0.8), Inches(1.65), card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. DEVELOPER VS BA MINDSET"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p = tf.add_paragraph()
    p.text = "• Pendekatan Developer (Teknis):\n  - Fokus struktur koding & query DB\n  - Implementasi tombol antarmuka UI\n  - 'Membuat tombol Confirm Putaway'\n  - Integrasi teknis endpoint REST API\n\n• Pendekatan BA (Logika Bisnis):\n  - Fokus logika bisnis & parameter input\n  - Validasi aturan, regulasi, & batas toleransi\n  - Analisis dampak ke operasional riil\n  - 'Sistem wajib memverifikasi berat barang ≤ kapasitas maksimum rak sebelum status lokasi diperbarui di database.'"
    p.font.size = Pt(9)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(8)

    create_card(s2, Inches(4.8), Inches(1.65), card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(5.0), Inches(1.85), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. RUMUS 3-LANGKAH ANALISIS BA"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf.add_paragraph()
    p.text = "Setiap fitur diubah menjadi spesifikasi bisnis melalui 3 langkah berurutan:\n\n1. FITUR / EVENT TRIGGER:\n   Tombol UI, pindaian sensor RFID, atau trigger waktu terjadwal.\n\n2. VALIDASI BUSINESS RULE:\n   Logika matematis, batasan hukum (SOLAS VGM, UU LLAJ), ambang toleransi, dan percabangan IF-THEN-ELSE.\n\n3. DAMPAK & PEMBARUAN SISTEM:\n   Perubahan status database, aksi perangkat keras fisik (palang buka/kunci), dan cetak dokumen resmi."
    p.font.size = Pt(9)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(8)

    create_card(s2, Inches(8.8), Inches(1.65), card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(9.0), Inches(1.85), card_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. CONTOH IMPLEMENTASI NYATA"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p = tf.add_paragraph()
    p.text = "Contoh Kasus 1: Fitur Putaway\n• Fitur: Tombol 'Confirm Putaway'\n• Rule: Cek Berat Palet ≤ Kapasitas Rak (1.500 kg/level) & Tinggi Palet ≤ 1.6m\n• Dampak: Jika lolos, status rak di DB jadi OCCUPIED; jika gagal, rak ditolak & dicarikan bin alternatif.\n\nContoh Kasus 2: Fitur Timbangan\n• Fitur: Sensor Loop Jembatan Timbang\n• Rule: Gross Weight ≤ 34.000 kg & Toleransi ±2% vs VGM SOLAS\n• Dampak: Lolos → palang buka & status APPROVED; Gagal → palang lock & status REJECTED."
    p.font.size = Pt(9)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(8)

    add_demo_banner(s2, "Konsep ini menjadi pondasi seluruh analisis fitur CIDP: Mengubah tombol menjadi kepatuhan operasional.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 3: MENJAWAB NOTULENSI: BUSINESS ROLE & EQUIPMENT ALLOCATION MATRIX
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Business Role & Matriks Alokasi Alat: Siapa yang Ngangkat, Nimbang, & Barangnya Apa?", web_page="dashboard.php?page=alat", web_url="http://localhost/.../alat.php")
    add_footer(s3, 3)

    t_left = Inches(0.8)
    t_top = Inches(1.65)
    t_width = Inches(11.7)
    t_height = Inches(4.55)

    shape_t = s3.shapes.add_table(6, 5, t_left, t_top, t_width, t_height)
    table = shape_t.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(2.5)
    table.columns[2].width = Inches(2.7)
    table.columns[3].width = Inches(2.7)
    table.columns[4].width = Inches(1.6)

    headers = ["Tahapan Operasi", "Karakteristik Barangnya", "Siapa yang Nimbang & Alat Sensor", "Siapa yang Ngangkat & Operator", "Modul Sistem"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_WHITE

    row_data = [
        ("1. Gerbang Masuk\n(Gate-In Weighbridge)", "Semua Muatan Truk:\n• FCL Dry Kontainer\n• Reefer Container\n• LCL Loose Cargo", "• Sensor WCS: Jembatan Timbang Pitless 80 Ton (Avery Load Cell)\n• Role: External Truck Driver & Tallyman Gate Inbound\n• Waktu: < 3 detik pembacaan", "Belum diangkat (masih di atas sasis truk trailer flatbed). Dilakukan verifikasi batas legal 34 ton.", "WCS + YMS\n(pages/gate.php)"),
        ("2. Lapangan Penumpukan\n(Yard Stacking)", "Kontainer 40 Feet Laden\n(Peti kemas isi muatan berat, bobot 20 - 32 Ton)", "Data berat diambil otomatis dari timbangan gate & dokumen VGM SOLAS yang tersimpan di database YMS.", "• SIAPA YANG NGANGKAT:\n  RTG Crane 45 Ton (Rubber Tyred Gantry) atau Reach Stacker 45 Ton.\n• OPERATOR: RTG Operator berlisensi SIO Kemnaker via VMT di kabin.", "YMS + WCS\n(pages/yard.php &\npages/alat.php)"),
        ("3. Penumpukan Kosong\n(Empty Stacking)", "Kontainer 20ft / 40ft Kosong\n(Empty Container, bobot tara 2,2 - 3,8 Ton)", "Ditimbang saat gate-in untuk verifikasi kondisi kosong (tare verification).", "• SIAPA YANG NGANGKAT:\n  Empty Container Handler (ECH) 10 Ton.\n• OPERATOR: ECH Operator berlisensi. Mampu menumpuk 6-8 tier.", "YMS + WCS\n(pages/alat.php)"),
        ("4. Gudang CFS\n(Stripping & Stuffing)", "Kargo Curah / LCL / Palet\n(Karton, drum kimia non-DG, suku cadang mesin)", "• Sensor WCS: Timbangan Palet Digital & DWS Mass Scanner di conveyor CFS.\n• Toleransi: ± 2%.", "• SIAPA YANG NGANGKAT:\n  Forklift 3-5 Ton (Diesel/Electric) & Reach Truck CFS.\n• OPERATOR: Forklift Driver & Tallyman CFS.", "WMS + WCS\n(pages/cfs.php &\nmaster_rak.php)"),
        ("5. Rel Siding KA 600M\n(Intermodal Rail)", "Kontainer 20ft & 40ft\n(Muatan intermodal ekspor/impor KAI Logistik)", "Timbangan dinamis gandar rel (Rail Weigh-in-Motion) batas maks 15 ton/gandar.", "• SIAPA YANG NGANGKAT:\n  Rail Mounted Gantry (RMG) Crane 45 Ton / Reach Stacker Intermodal.\n• OPERATOR: Intermodal Crane Operator.", "TMS + YMS\n(pages/intermodal.php)")
    ]

    for row_idx, r in enumerate(row_data, start=1):
        bg = C_CARD_BG if row_idx % 2 == 1 else RGBColor(241, 245, 249)
        for col_idx, text in enumerate(r):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(8)
            p.font.color.rgb = C_TEXT

    add_demo_banner(s3, "Jawaban Notulensi: Kontainer 40ft diangkat RTG / Reach Stacker 45T; Kosong oleh ECH; Palet CFS oleh Forklift 3-5T.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 4: FLOWCHART 1: MASTER END-TO-END DECISION FLOWCHART (GAMBAR RESMI)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Flowchart 1: Master End-to-End Decision Flowchart Terminal CIDP", web_page="dashboard.php?page=simulator", web_url="http://localhost/.../simulator.php")
    add_footer(s4, 4)

    e2e_img_path = os.path.join(base_dir, "FLOWCHART_END_TO_END_PROSES_BISNIS_CIDP.jpg")
    if os.path.exists(e2e_img_path):
        s4.shapes.add_picture(e2e_img_path, Inches(0.8), Inches(1.55), Inches(11.7), Inches(4.75))

    add_demo_banner(s4, "Flowchart End-to-End Resmi: Simbol Terminator, Decision Diamond, Process Box, & Cabang Ya/Tidak.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 5: FLOWCHART 2: DETAIL TIMBANGAN GATE-IN (GAMBAR RESMI)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Flowchart 2: Alur Keputusan Timbangan Dinamis Gate-In (WCS + YMS)", web_page="dashboard.php?page=gate", web_url="http://localhost/.../gate.php")
    add_footer(s5, 5)

    gate_img_path = os.path.join(base_dir, "FLOWCHART_TIMBANGAN_GATE_IN_DECISION.jpg")
    if os.path.exists(gate_img_path):
        s5.shapes.add_picture(gate_img_path, Inches(0.8), Inches(1.55), Inches(11.7), Inches(4.75))

    add_demo_banner(s5, "Menjawab Notulensi: Truk masuk timbangan, jika lulus ke mana vs jika tidak lulus belok ke Reject Lane.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 6: RINCIAN ANALITIS BUSINESS RULE 1: GATE-IN & TIMBANGAN DINAMIS
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Business Rule 1: Parameter Masukan & Formula Validasi Timbangan Gate-In", web_page="dashboard.php?page=gate", web_url="http://localhost/.../gate.php")
    add_footer(s6, 6)

    card_w2 = Inches(5.7)
    card_h2 = Inches(4.55)

    create_card(s6, Inches(0.8), Inches(1.65), card_w2, card_h2)
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. PARAMETER MASUKAN & FORMULA VALIDASI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p = tf.add_paragraph()
    p.text = "• Aktor & Hardware WCS:\n  - Driver Truk Eksternal, Gate Clerk / Tallyman\n  - Jembatan Timbang Pitless 80 Ton (Load Cell Avery Weigh-Tronix)\n  - Kamera ANPR (Plat Truk) & Kamera OCR Boks ISO 6346\n\n• Parameter Masukan (Input):\n  - Gross Weight (dari timbangan otomatis WCS)\n  - Tare Weight (berat kosong truk dari master data ERP)\n  - Net Weight = Gross Weight - Tare Weight\n  - Verified Gross Mass (VGM SOLAS) dari dokumen pengirim\n\n• Aturan Bisnis (Validation Rules):\n  1. Berat Kotor Total ≤ 34.000 kg (Kepatuhan UU No. 22/2009 Jalan Kelas 1)\n  2. Toleransi Selisih Berat ≤ ± 2.0% vs Manifest VGM SOLAS\n  3. Validasi Check Digit ISO 6346 Modulo 11"
    p.font.size = Pt(9)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s6, Inches(6.8), Inches(1.65), card_w2, card_h2)
    tb = s6.shapes.add_textbox(Inches(7.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. ALUR KEPUTUSAN OPERASIONAL (JIKA YA / JIKA TIDAK)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf.add_paragraph()
    p.text = "✅ JIKA LULUS TIMBANG (Valid & ≤ 34 Ton):\n• Palang barrier otomatis terbuka, lampu hijau aktif.\n• Status database berubah: GATE_IN_APPROVED.\n• Kiosk mencetak Container Movement Slip (CMS) ber-QR Code.\n• Ke mana tujuannya sesuai jenis barangnya:\n  - FCL Kering 40ft/20ft → Meluncur ke Blok Yard A/B/C/D.\n  - Reefer (Dingin) → Menuju Blok E (300 unit colokan 380V).\n  - Kargo Curah LCL → Menuju Gudang CFS Dock D1-D5.\n  - HazMat B3 → Menuju Blok F (Tanggul Isolasi Kimia).\n  - Pabean Jalur Merah → Menuju Hanggar Behandle Bea Cukai.\n\n❌ JIKA TIDAK LULUS (Overweight > 34 Ton / Selisih > 2%):\n• Palang barrier TETAP TERKUNCI RAPAT (Lock).\n• Lampu merah berkedip & sirine alarm aktif di pos gate.\n• Status database: OVERLOAD_REJECTED / HOLD.\n• Ke mana truk diarahkan?\n  - Truk DILARANG masuk yard (cegah jalan beton amblas).\n  - Masuk REJECT LANE / HOLDING BAY (Area Karantina Timbangan).\n  - Investigasi Supervisor: Rekoreksi dokumen atau U-Turn keluar."
    p.font.size = Pt(8.5)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    add_demo_banner(s6, "[LIVE DEMO]: Buka Halaman Web Gate (pages/gate.php) untuk simulasi timbangan lulus vs overweight.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 7: FLOWCHART 3: PENGANGKATAN & PENEMPATAN KONTAINER (GAMBAR RESMI)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Flowchart 3: Alokasi Alat Pengangkat & Stacking Rules 3D (YMS + WCS)", web_page="dashboard.php?page=yard", web_url="http://localhost/.../yard.php")
    add_footer(s7, 7)

    lift_img_path = os.path.join(base_dir, "FLOWCHART_PENGANGKATAN_DAN_STACKING_KONTAINER.jpg")
    if os.path.exists(lift_img_path):
        s7.shapes.add_picture(lift_img_path, Inches(0.8), Inches(1.55), Inches(11.7), Inches(4.75))

    add_demo_banner(s7, "Menjawab Notulensi: Kontainer 40ft diangkat RTG 45T, kosong oleh ECH, penempatan Bay Genap & Tier Berat Bawah.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 8: RINCIAN ANALITIS BUSINESS RULE 2: PENGANGKATAN & PENEMPATAN
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Business Rule 2: Pengangkatan & Penempatan Kontainer 40ft/20ft (Stacking Rules)", web_page="dashboard.php?page=yard", web_url="http://localhost/.../yard.php")
    add_footer(s8, 8)

    create_card(s8, Inches(0.8), Inches(1.65), card_w2, card_h2)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. ALOKASI ALAT PENGANGKAT BERDASARKAN SPESIFIKASI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p = tf.add_paragraph()
    p.text = "• Kontainer 40 Feet Heavy Laden (Muatan Penuh Berat):\n  - Alat Utama: RTG Crane (Rubber Tyred Gantry) 45 Ton.\n    Menjangkau rentang 6+1 baris dan 5+1 tinggi tier.\n  - Alat Sekunder: Reach Stacker 45 Ton (Kalmar / Konecranes).\n  - Operator: RTG/RS Operator bersertifikat SIO Kemnaker via VMT di kabin alat.\n\n• Kontainer 20 Feet Kosong (Empty Container):\n  - Alat: Empty Container Handler (ECH) 10 Ton.\n  - Operator: ECH Operator berlisensi.\n  - Kecepatan: Mampu menyusun hingga Tier 6-8 dengan cepat.\n\n• Kargo Curah / Palet di Gudang CFS:\n  - Alat: Forklift 3-5 Ton & Reach Truck CFS.\n  - Operator: CFS Forklift Operator & Tallyman.\n\n• Kontainer Intermodal di Rel Kereta Api 600M:\n  - Alat: RMG Crane 45 Ton / Reach Stacker Intermodal Leg."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s8, Inches(6.8), Inches(1.65), card_w2, card_h2)
    tb = s8.shapes.add_textbox(Inches(7.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. ATURAN PENEMPATAN 3D BAY-ROW-TIER (STACKING RULES)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf.add_paragraph()
    p.text = "• Struktur Hierarki Alamat: [Blok] - [Bay] - [Row] - [Tier]\n\n• Aturan Ukuran 20 Feet vs 40 Feet (Bay Rule):\n  - Bay Ganjil (01, 03, 05, dst.): Khusus 2 kontainer 20 Feet.\n  - Bay Genap (02, 06, 10, dst.): Khusus 1 kontainer 40 Feet.\n  - Larangan: Dilarang menaruh 40ft di atas dua 20ft jika posisi tidak seimbang atau tanpa twistlock pendukung.\n\n• Aturan Distribusi Berat (Weight Stacking Hierarchy):\n  - Tier 1 & 2 (Paling Bawah): Kontainer Heavy Laden (> 24 Ton) untuk menjaga kestabilan Center of Gravity.\n  - Tier 3 (Tengah): Kontainer Medium Weight (15 - 24 Ton).\n  - Tier 4 & 5 (Paling Atas): Kontainer Ringan (< 15 Ton) atau Kosong. Mencegah keruntuhan akibat tiupan angin.\n\n• Aturan Zonasi Karakteristik Kargo:\n  - Reefer: Blok E (dekat power plug 380V), maks Tier 3.\n  - HazMat B3: Blok F (jarak aman ≥ 50m, tanggul penampung)."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    add_demo_banner(s8, "[LIVE DEMO]: Buka Halaman Web Yard (pages/yard.php) untuk melihat visualisasi denah 3D Bay-Row-Tier.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 9: FITUR WMS INBOUND & INVENTORY MANAGEMENT (BAB 2 & 3 SILABUS)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Fitur WMS Inbound & Inventory Management: CFS & Cold Storage (Bab 2 & 3)", web_page="dashboard.php?page=cfs", web_url="http://localhost/.../cfs.php")
    add_footer(s9, 9)

    card_w3 = Inches(3.7)
    card_h3 = Inches(4.55)

    create_card(s9, Inches(0.8), Inches(1.65), card_w3, card_h3)
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w3 - Inches(0.4), card_h3 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. FITUR WMS INBOUND (BAB 2)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p = tf.add_paragraph()
    p.text = "• 2.1 ASN Matching:\n  Validasi otomatis nomor PO dari ERP. Jika kuantitas fisik > PO, picu status Over-receipt Alert.\n  Dampak: Menolak kelebihan barang tanpa otorisasi pembeli.\n\n• 2.2 Directed Putaway Engine:\n  Algoritma alokasi rak berdasarkan matriks kriteria: dimensi SKU, kategori HazMat, dan kebutuhan suhu ruang.\n  Dampak: Lokasi simpan optimal sesuai profil muatan.\n\n• 2.3 QC & Quarantine Status:\n  Otomatis mengubah status stok menjadi Blocked agar tidak teralokasi sebelum lolos inspeksi kualitas."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s9, Inches(4.8), Inches(1.65), card_w3, card_h3)
    tb = s9.shapes.add_textbox(Inches(5.0), Inches(1.85), card_w3 - Inches(0.4), card_h3 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. FITUR WMS INVENTORY (BAB 3)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf.add_paragraph()
    p.text = "• 3.1 Real-Time Bin Tracking:\n  Pemetaan hierarki Aisle-Rack-Shelf-Bin untuk memastikan keterlacakan (traceability) 100% dari palet hingga SKU individual.\n\n• 3.2 Automated Stock Rotation (FEFO / FIFO):\n  Mengunci alokasi barang berdasarkan Expiry Date paling dekat (FEFO) saat Pick Task dibuat oleh sistem.\n  Dampak: Menghilangkan risiko kedaluwarsa barang di gudang CFS.\n\n• 3.3 Cycle Counting Workflow:\n  Perhitungan stok berkala (blind counting) tanpa menghentikan operasional gudang, mencegah kecurangan operator."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s9, Inches(8.8), Inches(1.65), card_w3, card_h3)
    tb = s9.shapes.add_textbox(Inches(9.0), Inches(1.85), card_w3 - Inches(0.4), card_h3 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. COLD CHAIN & REEFER PLUGS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p = tf.add_paragraph()
    p.text = "• Manajemen Suhu Terkendali:\n  Chilled (2°C s/d 8°C) dan Frozen (-18°C).\n\n• Max Exposure Time (30 Menit):\n  Waktu maksimal kontainer reefer berada di luar colokan pendingin. YMS/WMS menghitung rute prioritas.\n\n• Telemetri Otomatis & Breach Alert:\n  Sensor digital mengirim suhu per 15 menit. Jika suhu naik melampaui threshold, sirine darurat berbunyi & teknisi genset dikerahkan otomatis.\n\n• Billing Kelistrikan 380V:\n  Perhitungan jam konsumsi daya presisi ke billing ERP."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    add_demo_banner(s9, "[LIVE DEMO]: Buka Halaman CFS (pages/cfs.php) & Reefer (pages/reefer.php) untuk menguji ASN & sensor suhu.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 10: FITUR WMS OUTBOUND & WCS HARDWARE CONTROL (BAB 4 & 5 SILABUS)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Fitur WMS Outbound & Kontrol Hardware WCS / WES Otomasi (Bab 4 & 5)", web_page="dashboard.php?page=scanner", web_url="http://localhost/.../scanner.php")
    add_footer(s10, 10)

    create_card(s10, Inches(0.8), Inches(1.65), card_w2, card_h2)
    tb = s10.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. FITUR WMS OUTBOUND (BAB 4)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p = tf.add_paragraph()
    p.text = "• 4.1 Slotting Optimization (Analisis ABC):\n  Penataan ulang lokasi simpan secara berkala berdasarkan frekuensi perputaran transaksi barang:\n  - Fast Moving (A): Di rak lantai bawah dekat pintu loading.\n  - Slow Moving (C): Di rak tingkat atas lorong belakang.\n  Dampak: Memangkas waktu putaway & retrieval hingga 35%.\n\n• 4.2 Picking Strategy Engine (Wave, Batch, Zone):\n  Penggabungan beberapa Sales Order menjadi satu jalur unduh (Wave) untuk memangkas jarak tempuh operator forklift.\n\n• 4.3 Packing & Weight Verification:\n  Pengecekan toleransi selisih berat (± 2%) pada timbangan otomatis sebelum pencetakan resi / shipping manifest."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s10, Inches(6.8), Inches(1.65), card_w2, card_h2)
    tb = s10.shapes.add_textbox(Inches(7.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. WCS HARDWARE CONTROL & ANOMALI (BAB 5)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf.add_paragraph()
    p.text = "• 5.1 Conveyor Diverter Control:\n  Penerjemahan sinyal scan barcode menjadi perintah fisik pusher/diverter untuk membelokkan paket ke jalur rute yang benar.\n\n• 5.2 Tunnel / Mass Scanner Integration (DWS):\n  Pencatatan data dimensi, berat (scale), dan GS1 scanner secara bersamaan dalam tempo < 500 ms.\n\n• 5.3 AMR / Crane Fleet Routing & Anti-Deadlock:\n  Manajemen jalur armada robotik & crane untuk mencegah tabrakan dan kondisi macet total (deadlock).\n\n• PENANGANAN ANOMALI WCS (EXCEPTION HANDLING):\n  - Kasus: Barcode rusak / No-Read Barcode pada conveyor.\n  - Aksi Sistem: WCS memerintahkan diverter membelokkan paket ke REJECT LANE, memicu Alert ke PDA operator untuk pindaian manual."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    add_demo_banner(s10, "[LIVE DEMO]: Buka Halaman Scanner (pages/scanner.php) untuk melihat integrasi DWS Mass Scanner & Diverter.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 11: FITUR YMS (YARD MANAGEMENT) & GATE MANAGEMENT (BAB 7 SILABUS)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Fitur YMS: Dock Scheduling, 3D Yard Tracking, & Demurrage Alert (Bab 7)", web_page="dashboard.php?page=yard", web_url="http://localhost/.../yard.php")
    add_footer(s11, 11)

    create_card(s11, Inches(0.8), Inches(1.65), card_w3, card_h3)
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w3 - Inches(0.4), card_h3 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "7.1 DOCK DOOR SCHEDULING"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p = tf.add_paragraph()
    p.text = "• Penjadwalan Time-Slot Armada:\n  Sistem menjadwalkan kedatangan truk eksternal berdasarkan slot waktu 30 menit (Appointment Window).\n\n• Tujuan Bisnis:\n  Mencegah penumpukan antrean truk trailer di jalan akses utama dry port.\n\n• Integrasi Cold Dock Seal:\n  Mengalokasikan dock door khusus yang memiliki kantong kedap udara (inflatable seal) agar hawa dingin gudang tidak terbuang saat truk loading."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s11, Inches(4.8), Inches(1.65), card_w3, card_h3)
    tb = s11.shapes.add_textbox(Inches(5.0), Inches(1.85), card_w3 - Inches(0.4), card_h3 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "7.2 TRAILER POSITION TRACKING"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf.add_paragraph()
    p.text = "• Visual Yard Map Real-Time:\n  Pemetaan grafis seluruh blok lapangan penumpukan 35 Ha.\n\n• Status Peti Kemas:\n  - EMPTY (Kosong)\n  - LADEN (Terisi Penuh)\n  - IN_STRIPPING (Bongkar Muat)\n  - CUSTOMS_HOLD (Karantina Bea Cukai)\n\n• Keterlacakan Posisi:\n  Setiap pergerakan oleh RTG Crane atau Reach Stacker langsung memperbarui database YMS via integrasi VMT WiFi 6 terminal."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s11, Inches(8.8), Inches(1.65), card_w3, card_h3)
    tb = s11.shapes.add_textbox(Inches(9.0), Inches(1.85), card_w3 - Inches(0.4), card_h3 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "7.3 GATE & DEMURRAGE ALERT"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD
    p = tf.add_paragraph()
    p.text = "• Otomasi Pencatatan Gate:\n  Timestamp gate-in dan gate-out tercatat otomatis presisi ke milidetik.\n\n• Dwelling Time Rule:\n  - Batas Bebas (Free Time) = 3 Hari.\n  - Hari ke-4 s/d 7 = Rp 250.000/hari.\n  - Hari ke-8 ke atas = Rp 500.000/hari.\n\n• Alert Notifikasi Otomatis:\n  Sistem mengirim peringatan dini (Early Warning) ke Shipper pada Hari ke-2 bahwa free time akan berakhir, mendorong percepatan pengambilan barang."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    add_demo_banner(s11, "[LIVE DEMO]: Buka Halaman Billing (pages/billing.php) untuk melihat simulasi kalkulasi demurrage progresif.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 12: FITUR TMS TRANSPORTATION & REL SIDING KA (BAB 6 & 8 SILABUS)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Fitur TMS Transportation & Rel Siding Kereta Api Intermodal 600M (Bab 6 & 8)", web_page="dashboard.php?page=trucking", web_url="http://localhost/.../trucking.php")
    add_footer(s12, 12)

    create_card(s12, Inches(0.8), Inches(1.65), card_w2, card_h2)
    tb = s12.shapes.add_textbox(Inches(1.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. FITUR TMS PLANNING & EXECUTION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p = tf.add_paragraph()
    p.text = "• 6.1 Dynamic Route Optimization Engine:\n  Algoritma rute terpendek dengan memperhitungkan kemacetan, jam operasional ganjil-genap, dan kelas jalan tonase.\n\n• 6.2 3D Load Planning & Consolidation:\n  Simulasi tata letak kubikasi kargo di bak truk untuk memaksimalkan Cubic Utilization hingga 92%.\n\n• 6.3 Carrier Rate Audit & Settlement:\n  Audit otomatis biaya pengiriman vendor pihak ketiga berdasarkan matriks tarif jarak dan zona wilayah.\n\n• 8.1 Real-Time GPS Tracking & Geofencing:\n  Pemantauan rute armada truk. Notifikasi instan jika supir keluar rute resmi atau berhenti > 45 menit tanpa izin.\n\n• 8.2 Driver Mobile App & e-POD:\n  Bukti serah terima digital: Foto barang, TTD digital, & GPS timestamp."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    create_card(s12, Inches(6.8), Inches(1.65), card_w2, card_h2)
    tb = s12.shapes.add_textbox(Inches(7.0), Inches(1.85), card_w2 - Inches(0.4), card_h2 - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. REL SIDING KERETA API 600M (INTERMODAL)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf.add_paragraph()
    p.text = "• Kapasitas Rel Siding CIDP:\n  Panjang jalur 600 meter melayani 1 rangkaian lokomotif KAI Logistik bermuatan 30 gerbong datar (flatcar) atau 60 TEUs.\n\n• Sinkronisasi Transshipment Kereta Api:\n  - TMS & YMS bertukar data manifest KA sebelum kedatangan.\n  - Crane Pengangkat: RMG Crane 45 Ton / Reach Stacker Intermodal.\n  - Waktu Bongkar Muat: Target < 4 menit per peti kemas.\n\n• Keunggulan Moda Kereta Api vs Truk Jalan Raya:\n  - Efisiensi Biaya Bahan Bakar: Menghemat hingga 40% vs trucking jalan raya.\n  - Bebas Macet Tol Cikampek / Pantura: Kepastian waktu tiba (On-Time Delivery > 98%).\n  - Ramah Lingkungan (Green Logistics): Menurunkan jejak karbon terminal hingga 65%."
    p.font.size = Pt(8.8)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    add_demo_banner(s12, "[LIVE DEMO]: Buka Halaman Intermodal (pages/intermodal.php) & Trucking (pages/trucking.php).", top=Inches(6.40))

    # =========================================================================
    # SLIDE 13: FLOWCHART 4: ALUR INTEGRASI 5 SISTEM LOGISTIK (BAB 12, 13, 14)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Flowchart 4: Alur Integrasi 5 Modul Logistik Terpadu (Bab 12, 13, & 14)", web_page="dashboard.php", web_url="http://localhost/.../dashboard.php")
    add_footer(s13, 13)

    integ_img_path = os.path.join(base_dir, "FLOWCHART_INTEGRASI_5_MODUL_LOGISTIK.jpg")
    if os.path.exists(integ_img_path):
        s13.shapes.add_picture(integ_img_path, Inches(0.8), Inches(1.55), Inches(11.7), Inches(4.75))

    add_demo_banner(s13, "Standar Silabus Bab 12 & 13: ERP (SO) -> WMS (Cold Zone) -> WCS (AS/RS Scale) -> YMS (Dock) -> TMS (Fleet) -> Invoicing.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 14: SOLUSI KEBERSIHAN: DI MANA DITARUHNYA DI APLIKASI WEB CIDP
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Menjawab Dilemma: Di Mana Ditaruhnya & Konsepnya di Aplikasi Web CIDP?", web_page="dashboard.php", web_url="http://localhost/.../dashboard.php")
    add_footer(s14, 14)

    app_img_path = os.path.join(base_dir, "PETA_ARSITEKTUR_DAN_KONSEP_APLIKASI_CIDP.jpg")
    if os.path.exists(app_img_path):
        s14.shapes.add_picture(app_img_path, Inches(0.8), Inches(1.55), Inches(11.7), Inches(4.75))

    add_demo_banner(s14, "Solusi Kompleksitas: Pemisahan 5 Modul Terstruktur (WCS, YMS, WMS, TMS, ERP) menjaga kode rapi dan modular.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 15: MATRIKS INTEGRASI DATA FITUR & EXCEPTION HANDLING (BAB 9, 10, 11)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "Matriks Integrasi Data Fitur (Data Mapping), Anomali, & FSD (Bab 9, 10, 11)", web_page="dashboard.php?page=settings", web_url="http://localhost/.../settings.php")
    add_footer(s15, 15)

    shape_t2 = s15.shapes.add_table(5, 4, Inches(0.8), Inches(1.65), Inches(7.5), Inches(4.55))
    table2 = shape_t2.table
    table2.columns[0].width = Inches(1.8)
    table2.columns[1].width = Inches(1.8)
    table2.columns[2].width = Inches(2.1)
    table2.columns[3].width = Inches(1.8)

    headers2 = ["Modul Fitur", "Data Input (Source)", "Proses Validasi Sistem", "Data Output (Target)"]
    for i, h in enumerate(headers2):
        cell = table2.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(9)
        p.font.color.rgb = C_WHITE

    row_data2 = [
        ("ASN Receiving\n(WMS)", "PO Data dari ERP &\nDokumen Delivery", "Cek kesesuaian nomor PO & toleransi kuantitas", "Status: Received /\nOver-receipt Alert"),
        ("Picking Engine\n(WMS / WCS)", "Order List dari ERP &\nLokasi Bin CFS", "Cek aturan FEFO & kapasitas beban rak", "Task Command ke\nPDA / AMR / Forklift"),
        ("Mass Scanning\n(WCS)", "Barcode & Weight Sensor\nTunnel DWS (<500ms)", "Cek kecocokan manifes & toleransi berat ±2%", "Signal Diverter &\nResi Label Tercetak"),
        ("Route Planner\n(TMS)", "Alamat Kirim, Bobot,\n& Batasan Jam Jalan", "Algoritma rute tercepat & optimasi tonase", "Manifest Jalan &\nGPS Track Assignment")
    ]

    for row_idx, r in enumerate(row_data2, start=1):
        bg = C_CARD_BG if row_idx % 2 == 1 else RGBColor(241, 245, 249)
        for col_idx, text in enumerate(r):
            cell = table2.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(8)
            p.font.color.rgb = C_TEXT

    create_card(s15, Inches(8.5), Inches(1.65), Inches(4.0), Inches(4.55))
    tb = s15.shapes.add_textbox(Inches(8.7), Inches(1.85), Inches(3.6), Inches(4.15))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PENANGANAN ANOMALI & FSD STANDAR"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_RED
    p = tf.add_paragraph()
    p.text = "1. Anomali Overweight di Gate Timbangan:\n• Gejala: Berat > 34 ton atau selisih > 2%.\n• Penanganan: Palang lock, truk dialihkan ke Holding Bay Karantina, investigasi dokumen.\n\n2. No-Read Barcode pada Conveyor WCS:\n• Gejala: Barcode robek / scanner gagal baca.\n• Penanganan: Diverter belokkan paket ke Reject Lane + alert PDA operator.\n\n3. Shortage Stok saat Picking WMS:\n• Gejala: Fisik barang di rak kurang dari order.\n• Penanganan: Klik 'Item Shortage', sistem picu Blind Cycle Count & alokasi rak alternatif.\n\n4. 4 Pilar FSD Standar Silabus:\n• User Role & Permission\n• Trigger Condition\n• Input Fields & Validation Rules\n• System Response & Error Handling"
    p.font.size = Pt(8.5)
    p.font.color.rgb = C_TEXT
    p.space_before = Pt(6)

    add_demo_banner(s15, "Matriks integrasi dan skenario exception handling menjamin operasional stabil 24/7 tanpa henti.", top=Inches(6.40))

    # =========================================================================
    # SLIDE 16: KESIMPULAN REKOMENDASI KONSULTAN & ROADMAP MENUJU GO-LIVE
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    bg16 = s16.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg16.fill.solid()
    bg16.fill.fore_color.rgb = C_NAVY
    bg16.line.fill.background()

    tbox16 = s16.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
    tf16 = tbox16.text_frame
    tf16.word_wrap = True
    p = tf16.paragraphs[0]
    p.text = "KESIMPULAN REKOMENDASI KONSULTAN & ROADMAP MENUJU GO-LIVE"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p = tf16.add_paragraph()
    p.text = "STRATEGI EKSEKUTIF MENUJU IMPLEMENTASI TERMINAL INTERMODAL DRY PORT KELAS DUNIA"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p.space_before = Pt(3)

    col_w_fin = Inches(3.7)
    card_h_fin = Inches(4.6)
    top_fin = Inches(1.85)

    c16_1 = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_fin, col_w_fin, card_h_fin)
    c16_1.fill.solid()
    c16_1.fill.fore_color.rgb = RGBColor(15, 33, 64)
    c16_1.line.color.rgb = RGBColor(30, 58, 110)
    tf16_1 = c16_1.text_frame
    tf16_1.word_wrap = True
    p = tf16_1.paragraphs[0]
    p.text = "1. KEPATUHAN STANDAR REGULASI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p = tf16_1.add_paragraph()
    p.text = "• Verifikasi Berat Internasional:\n  Implementasi jembatan timbang 80T patuh 100% pada SOLAS VGM Bab VI & UU RI No. 22/2009.\n\n• Integritas Penomoran ISO 6346:\n  Validasi Modulo 11 check digit mencegah kesalahan identitas peti kemas.\n\n• Kepatuhan Bea Cukai CEISA 4.0:\n  Kunci otomatis (Customs Lock) memastikan barang tanpa SPPB tidak bocor ke luar pelabuhan kering."
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(8)

    c16_2 = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), top_fin, col_w_fin, card_h_fin)
    c16_2.fill.solid()
    c16_2.fill.fore_color.rgb = RGBColor(15, 33, 64)
    c16_2.line.color.rgb = RGBColor(30, 58, 110)
    tf16_2 = c16_2.text_frame
    tf16_2.word_wrap = True
    p = tf16_2.paragraphs[0]
    p.text = "2. EFISIENSI & ZERO LEAKAGE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(52, 211, 153)
    p = tf16_2.add_paragraph()
    p.text = "• Zero Revenue Leakage:\n  Penguncian palang barrier gate-out hingga invoice berstatus PAID menghilangkan potensi kebocoran pendapatan.\n\n• Pangkas Waktu Gate In:\n  Otomasi ANPR, OCR, dan timbangan memangkas waktu proses dari 15 menit menjadi 90 detik per truk.\n\n• Penghematan Logistik Intermodal:\n  Konektivitas rel KA 600m memangkas biaya transportasi jarak jauh hingga 40%."
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(8)

    c16_3 = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), top_fin, col_w_fin, card_h_fin)
    c16_3.fill.solid()
    c16_3.fill.fore_color.rgb = RGBColor(15, 33, 64)
    c16_3.line.color.rgb = RGBColor(30, 58, 110)
    tf16_3 = c16_3.text_frame
    tf16_3.word_wrap = True
    p = tf16_3.paragraphs[0]
    p.text = "3. ROADMAP MENUJU GO-LIVE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(251, 191, 36)
    p = tf16_3.add_paragraph()
    p.text = "• FASE 1 (Bulan 1-3): Fondasi Sistem\n  Finalisasi modul Web YMS, migrasi skema MySQL, dan integrasi modul Odoo ERP.\n\n• FASE 2 (Bulan 4-6): Kalibrasi IoT\n  Pemasangan sensor timbangan Avery, kamera ANPR/OCR, dan commissioning VMT RTG.\n\n• FASE 3 (Bulan 7-9): Uji Coba Lapangan\n  Integrasi API CEISA 4.0 Bea Cukai, simulasi KA Logistik, dan pelatihan operator terminal."
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_before = Pt(8)

    f16 = s16.shapes.add_textbox(Inches(0.8), Inches(6.75), Inches(11.7), Inches(0.4))
    tf16_f = f16.text_frame
    p = tf16_f.paragraphs[0]
    p.text = "Conclusion Supply Chain Consultant  •  Terima Kasih  •  Sesi Tanya Jawab & Diskusi Terbuka"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.alignment = PP_ALIGN.CENTER

    out_pptx1 = os.path.join(base_dir, "PRESENTASI_BUSINESS_ANALYST_CIDP.pptx")
    prs.save(out_pptx1)
    print("Saved:", out_pptx1)

    out_pptx2 = os.path.join(base_dir, "Conclusion Intermodal Dry Port_Business Analysis.pptx")
    prs.save(out_pptx2)
    print("Saved:", out_pptx2)

    root_ba_dir = r"C:\xampp\htdocs\Conclusion_Intermodal_Dry_Port\business_analyst"
    if os.path.exists(root_ba_dir):
        out_pptx3 = os.path.join(root_ba_dir, "PRESENTASI_BUSINESS_ANALYST_CIDP.pptx")
        prs.save(out_pptx3)
        print("Saved:", out_pptx3)

if __name__ == "__main__":
    create_ba_presentation()
